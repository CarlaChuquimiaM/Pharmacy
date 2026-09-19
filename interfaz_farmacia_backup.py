import sqlite3
from datetime import datetime
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
from tkinter import filedialog
from openpyxl import load_workbook
from PIL import Image, ImageTk
import os
def conectar():
    conexion = sqlite3.connect("farmacia.db")
    cursor = conexion.cursor()

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS productos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo TEXT UNIQUE,
        nombre TEXT,
        marca TEXT,
        precio REAL,
        stock INTEGER
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS ventas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre_producto TEXT,
        cantidad INTEGER,
        total REAL,
        fecha TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS promociones (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        producto_id INTEGER NOT NULL,
        descuento REAL NOT NULL,
        fecha_inicio TEXT NOT NULL,
        fecha_fin TEXT NOT NULL,
        activa INTEGER DEFAULT 1
    )
    """)

    cursor.execute("PRAGMA table_info(productos)")
    columnas = [columna[1] for columna in cursor.fetchall()]

    if "vencimiento" not in columnas:
        cursor.execute("ALTER TABLE productos ADD COLUMN vencimiento TEXT")

    conexion.commit()
    return conexion, cursor
def agregar_producto():

    ventana_agregar = tk.Toplevel()
    ventana_agregar.title("Agregar producto")
    ventana_agregar.geometry("350x380")

    tk.Label(ventana_agregar, text="Código:").pack(pady=5)
    entry_codigo = tk.Entry(ventana_agregar, width=30)
    entry_codigo.pack()

    tk.Label(ventana_agregar, text="Producto:").pack(pady=5)
    entry_nombre = tk.Entry(ventana_agregar, width=30)
    entry_nombre.pack()

    tk.Label(ventana_agregar, text="Marca:").pack(pady=5)
    entry_marca = tk.Entry(ventana_agregar, width=30)
    entry_marca.pack()

    tk.Label(ventana_agregar, text="Precio:").pack(pady=5)
    entry_precio = tk.Entry(ventana_agregar, width=30)
    entry_precio.pack()

    tk.Label(ventana_agregar, text="Stock:").pack(pady=5)
    entry_stock = tk.Entry(ventana_agregar, width=30)
    entry_stock.pack()

    tk.Label(ventana_agregar, text="Vencimiento (YYYY-MM-DD):").pack(pady=5)
    entry_vencimiento = tk.Entry(ventana_agregar, width=30)
    entry_vencimiento.pack()

    def guardar():
        codigo = entry_codigo.get().strip()
        nombre = entry_nombre.get().strip()
        marca = entry_marca.get().strip()
        vencimiento = entry_vencimiento.get().strip()

        try:
            precio = float(entry_precio.get())
            stock = int(entry_stock.get())
        except:
            messagebox.showerror("Error", "Precio o stock inválido")
            return

        if nombre == "":
            messagebox.showwarning("Aviso", "Ingresa un producto")
            return

        if vencimiento != "":
            try:
                datetime.strptime(vencimiento, "%Y-%m-%d")
            except:
                messagebox.showerror("Error", "La fecha debe ir como YYYY-MM-DD")
                return

        conexion, cursor = conectar()
        cursor.execute(
            "INSERT INTO productos (codigo, nombre, marca, precio, stock, vencimiento) VALUES (?, ?, ?, ?, ?, ?)",
            (codigo, nombre, marca, precio, stock, vencimiento)
        )
        conexion.commit()
        conexion.close()

        messagebox.showinfo("Éxito", "Producto agregado")
        ventana_agregar.destroy()

    tk.Button(
        ventana_agregar,
        text="Guardar producto",
        bg="#27ae60",
        fg="white",
        command=guardar
    ).pack(pady=15)
def ver_inventario():

    def cargar_productos(filtro=""):
        for fila_tabla in tabla.get_children():
            tabla.delete(fila_tabla)

        conexion, cursor = conectar()

        if filtro:
            cursor.execute(
                "SELECT * FROM productos WHERE nombre LIKE ? OR codigo LIKE ?",
                ("%" + filtro + "%", "%" + filtro + "%")
            )
        else:
            cursor.execute("SELECT * FROM productos")

        productos = cursor.fetchall()
        conexion.close()

        hoy = datetime.now().date()

        for producto in productos:
            stock = producto[5]
            vencimiento = producto[6] if len(producto) > 6 else ""

            estado = "Normal"

            if stock <= 10:
                estado = "Stock bajo"

            if vencimiento:
                try:
                    fecha_venc = datetime.strptime(vencimiento, "%Y-%m-%d").date()
                    dias = (fecha_venc - hoy).days

                    if dias < 0:
                        estado = "Vencido"
                    elif dias <= 30:
                        estado = "Por vencer"
                except:
                    pass

            tags = ()
            if estado == "Stock bajo":
                tags = ("bajo",)
            elif estado == "Por vencer":
                tags = ("por_vencer",)
            elif estado == "Vencido":
                tags = ("vencido",)

            tabla.insert(
                "",
                "end",
                values=(
                    producto[0],
                    producto[1],
                    producto[2],
                    producto[3],
                    producto[4],
                    producto[5],
                    vencimiento,
                    estado
                ),
                tags=tags
            )

    def buscar_producto():
        texto_busqueda = entry_buscar.get().strip()
        cargar_productos(texto_busqueda)

    def editar_producto():
        seleccionado = tabla.selection()

        if not seleccionado:
            messagebox.showwarning("Aviso", "Selecciona un producto")
            return

        datos = tabla.item(seleccionado[0], "values")
        id_producto = datos[0]
        precio_actual = datos[4]
        stock_actual = datos[5]
        vencimiento_actual = datos[6]

        ventana_editar = tk.Toplevel()
        ventana_editar.title("Editar producto")
        ventana_editar.geometry("300x260")

        tk.Label(ventana_editar, text="Precio:").pack(pady=5)
        entry_precio = tk.Entry(ventana_editar)
        entry_precio.pack()
        entry_precio.insert(0, precio_actual)

        tk.Label(ventana_editar, text="Stock:").pack(pady=5)
        entry_stock = tk.Entry(ventana_editar)
        entry_stock.pack()
        entry_stock.insert(0, stock_actual)

        tk.Label(ventana_editar, text="Vencimiento (YYYY-MM-DD):").pack(pady=5)
        entry_vencimiento = tk.Entry(ventana_editar)
        entry_vencimiento.pack()
        entry_vencimiento.insert(0, vencimiento_actual)

        def guardar():
            try:
                nuevo_precio = float(entry_precio.get())
                nuevo_stock = int(entry_stock.get())
            except:
                messagebox.showerror("Error", "Datos inválidos")
                return

            nuevo_vencimiento = entry_vencimiento.get().strip()

            if nuevo_vencimiento != "":
                try:
                    datetime.strptime(nuevo_vencimiento, "%Y-%m-%d")
                except:
                    messagebox.showerror("Error", "La fecha debe ir como YYYY-MM-DD")
                    return

            conexion, cursor = conectar()
            cursor.execute(
                "UPDATE productos SET precio = ?, stock = ?, vencimiento = ? WHERE id = ?",
                (nuevo_precio, nuevo_stock, nuevo_vencimiento, id_producto)
            )
            conexion.commit()
            conexion.close()

            messagebox.showinfo("Éxito", "Producto actualizado")
            ventana_editar.destroy()
            buscar_producto()

        tk.Button(
            ventana_editar,
            text="Guardar cambios",
            bg="#27ae60",
            fg="white",
            command=guardar
        ).pack(pady=15)

    def eliminar_producto():
        seleccionado = tabla.selection()

        if not seleccionado:
            messagebox.showwarning("Aviso", "Selecciona un producto")
            return

        datos = tabla.item(seleccionado[0], "values")
        id_producto = datos[0]

        confirmar = messagebox.askyesno("Confirmar", "¿Eliminar este producto?")

        if not confirmar:
            return

        conexion, cursor = conectar()
        cursor.execute("DELETE FROM productos WHERE id = ?", (id_producto,))
        conexion.commit()
        conexion.close()

        messagebox.showinfo("Éxito", "Producto eliminado")
        buscar_producto()

    def importar_excel():
        archivo = filedialog.askopenfilename(
            title="Seleccionar archivo Excel",
            filetypes=[("Archivos Excel", "*.xlsx")]
        )

        if not archivo:
            return

        try:
            libro = load_workbook(archivo, data_only=True)

            if "INVENTARIO" in libro.sheetnames:
                hoja = libro["INVENTARIO"]
            else:
                hoja = libro.active

            conexion, cursor = conectar()

            importados = 0
            actualizados = 0
            omitidos = 0

            for fila_excel in hoja.iter_rows(min_row=2, values_only=True):
                if fila_excel is None:
                    continue

                if len(fila_excel) < 7:
                    omitidos += 1
                    continue

                codigo = fila_excel[0]
                nombre = fila_excel[1]
                marca = fila_excel[3]
                entrada = fila_excel[4]
                salida = fila_excel[5]
                stock_excel = fila_excel[6]

                if nombre is None:
                    omitidos += 1
                    continue

                codigo = "" if codigo is None else str(codigo).strip()
                nombre = str(nombre).strip()
                marca = "" if marca is None else str(marca).strip()

                try:
                    entrada = 0 if entrada is None else int(entrada)
                except:
                    entrada = 0

                try:
                    salida = 0 if salida is None else int(salida)
                except:
                    salida = 0

                try:
                    if stock_excel is None:
                        stock = entrada - salida
                    else:
                        stock = int(stock_excel)
                except:
                    stock = 0

                if stock < 0:
                    stock = 0

                precio = 0.0
                vencimiento = ""

                cursor.execute("SELECT id FROM productos WHERE codigo = ?", (codigo,))
                existente = cursor.fetchone()

                if existente:
                    cursor.execute(
                        "UPDATE productos SET nombre = ?, marca = ?, stock = ? WHERE codigo = ?",
                        (nombre, marca, stock, codigo)
                    )
                    actualizados += 1
                else:
                    cursor.execute(
                        "INSERT INTO productos (codigo, nombre, marca, precio, stock, vencimiento) VALUES (?, ?, ?, ?, ?, ?)",
                        (codigo, nombre, marca, precio, stock, vencimiento)
                    )
                    importados += 1

            conexion.commit()
            conexion.close()

            messagebox.showinfo(
                "Importación completada",
                f"Productos nuevos: {importados}\n"
                f"Productos actualizados: {actualizados}\n"
                f"Filas omitidas: {omitidos}"
            )

            cargar_productos(entry_buscar.get().strip())

        except Exception as e:
            messagebox.showerror("Error", f"No se pudo importar el Excel:\n{e}")

    ventana = tk.Toplevel()
    ventana.title("Inventario")
    ventana.geometry("1000x520")

    frame_busqueda = tk.Frame(ventana)
    frame_busqueda.pack(pady=10)

    tk.Label(frame_busqueda, text="Buscar medicamento:").grid(row=0, column=0, padx=5)

    entry_buscar = tk.Entry(frame_busqueda, width=30)
    entry_buscar.grid(row=0, column=1, padx=5)

    btn_buscar = tk.Button(
        frame_busqueda,
        text="Buscar",
        bg="#3498db",
        fg="white",
        command=buscar_producto
    )
    btn_buscar.grid(row=0, column=2, padx=5)

    btn_mostrar_todos = tk.Button(
        frame_busqueda,
        text="Mostrar todos",
        bg="#7f8c8d",
        fg="white",
        command=lambda: [entry_buscar.delete(0, tk.END), cargar_productos()]
    )
    btn_mostrar_todos.grid(row=0, column=3, padx=5)

    columnas = ("ID", "Codigo", "Producto", "Marca", "Precio", "Stock", "Vencimiento", "Estado")

    tabla = ttk.Treeview(ventana, columns=columnas, show="headings")

    tabla.heading("ID", text="ID")
    tabla.heading("Codigo", text="Código")
    tabla.heading("Producto", text="Producto")
    tabla.heading("Marca", text="Marca")
    tabla.heading("Precio", text="Precio")
    tabla.heading("Stock", text="Stock")
    tabla.heading("Vencimiento", text="Vencimiento")
    tabla.heading("Estado", text="Estado")

    tabla.column("ID", width=50)
    tabla.column("Codigo", width=100)
    tabla.column("Producto", width=220)
    tabla.column("Marca", width=140)
    tabla.column("Precio", width=90)
    tabla.column("Stock", width=70)
    tabla.column("Vencimiento", width=110)
    tabla.column("Estado", width=120)

    tabla.pack(fill="both", expand=True, pady=10)

    tabla.tag_configure("bajo", background="#ffb3b3")
    tabla.tag_configure("por_vencer", background="#ffe599")
    tabla.tag_configure("vencido", background="#ff9999")

    frame_botones = tk.Frame(ventana)
    frame_botones.pack(pady=10)

    btn_editar = tk.Button(
        frame_botones,
        text="Editar producto",
        width=16,
        bg="#3498db",
        fg="white",
        command=editar_producto
    )
    btn_editar.grid(row=0, column=0, padx=8)

    btn_eliminar = tk.Button(
        frame_botones,
        text="Eliminar producto",
        width=16,
        bg="#e74c3c",
        fg="white",
        command=eliminar_producto
    )
    btn_eliminar.grid(row=0, column=1, padx=8)

    btn_actualizar = tk.Button(
        frame_botones,
        text="Actualizar inventario",
        width=16,
        bg="#27ae60",
        fg="white",
        command=lambda: cargar_productos(entry_buscar.get().strip())
    )
    btn_actualizar.grid(row=0, column=2, padx=8)

    btn_agregar = tk.Button(
        frame_botones,
        text="Agregar producto",
        width=16,
        bg="#2980b9",
        fg="white",
        command=agregar_producto
    )
    btn_agregar.grid(row=0, column=3, padx=8)

    btn_importar = tk.Button(
        frame_botones,
        text="Importar Excel",
        width=16,
        bg="#16a085",
        fg="white",
        command=importar_excel
    )
    btn_importar.grid(row=0, column=4, padx=8)

    cargar_productos()
def vender_producto():
    conexion, cursor = conectar()
    cursor.execute("SELECT nombre FROM productos ORDER BY nombre")
    lista_productos = [fila[0] for fila in cursor.fetchall()]
    conexion.close()

    carrito = []

    def filtrar_productos(event=None):
        texto = combo_nombre.get().strip().lower()

        if texto == "":
            combo_nombre["values"] = lista_productos
        else:
            filtrados = [p for p in lista_productos if texto in p.lower()]
            combo_nombre["values"] = filtrados

    def obtener_precio_con_promocion(producto_id, precio_normal):
        hoy = datetime.now().date()

        try:
            conexion, cursor = conectar()
            cursor.execute("""
                SELECT descuento, fecha_inicio, fecha_fin
                FROM promociones
                WHERE producto_id = ? AND activa = 1
                ORDER BY id DESC
            """, (producto_id,))
            promociones = cursor.fetchall()
            conexion.close()
        except:
            return precio_normal, 0

        for promo in promociones:
            descuento, fecha_inicio, fecha_fin = promo
            try:
                inicio = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
                fin = datetime.strptime(fecha_fin, "%Y-%m-%d").date()

                if inicio <= hoy <= fin:
                    precio_final = precio_normal - (precio_normal * (descuento / 100))
                    return round(precio_final, 2), descuento
            except:
                pass

        return precio_normal, 0

    def buscar_producto():
        nombre_producto = combo_nombre.get().strip()

        if not nombre_producto:
            messagebox.showwarning("Aviso", "Selecciona o escribe un medicamento")
            return

        conexion, cursor = conectar()
        cursor.execute("SELECT * FROM productos WHERE nombre = ?", (nombre_producto,))
        producto = cursor.fetchone()
        conexion.close()

        if producto is None:
            messagebox.showerror("Error", "Producto no encontrado")
            label_info.config(text="Medicamento no encontrado", fg="red")
            return

        producto_id = producto[0]
        precio = producto[4]
        stock = producto[5]

        precio_final, descuento = obtener_precio_con_promocion(producto_id, precio)

        if descuento > 0:
            label_info.config(
                text=f"Precio normal: {precio} Bs\nDescuento: {descuento}%\nPrecio final: {precio_final} Bs\nStock disponible: {stock}",
                fg="#1f3b5c"
            )
        else:
            label_info.config(
                text=f"Precio: {precio} Bs   |   Stock disponible: {stock}",
                fg="#1f3b5c"
            )

    def actualizar_tabla_carrito():
        for fila in tabla_carrito.get_children():
            tabla_carrito.delete(fila)

        total_general = 0

        for i, item in enumerate(carrito):
            subtotal = item["subtotal"]
            total_general += subtotal

            tabla_carrito.insert(
                "",
                "end",
                iid=str(i),
                values=(
                    item["nombre"],
                    item["cantidad"],
                    item["precio_normal"],
                    f'{item["descuento"]}%',
                    item["precio_final"],
                    subtotal
                )
            )

        label_total.config(text=f"TOTAL GENERAL: {round(total_general, 2)} Bs")

    def agregar_al_carrito():
        nombre_producto = combo_nombre.get().strip()

        if not nombre_producto:
            messagebox.showwarning("Aviso", "Selecciona o escribe un medicamento")
            return

        try:
            cantidad_vendida = int(entry_cantidad.get().strip())
            if cantidad_vendida <= 0:
                messagebox.showerror("Error", "La cantidad debe ser mayor a 0")
                return
        except:
            messagebox.showerror("Error", "Cantidad no válida")
            return

        conexion, cursor = conectar()
        cursor.execute("SELECT * FROM productos WHERE nombre = ?", (nombre_producto,))
        producto = cursor.fetchone()
        conexion.close()

        if producto is None:
            messagebox.showerror("Error", "Producto no encontrado")
            return

        producto_id = producto[0]
        precio_normal = producto[4]
        stock_actual = producto[5]

        precio_final, descuento = obtener_precio_con_promocion(producto_id, precio_normal)

        cantidad_en_carrito = 0
        for item in carrito:
            if item["producto_id"] == producto_id:
                cantidad_en_carrito += item["cantidad"]

        if cantidad_vendida + cantidad_en_carrito > stock_actual:
            messagebox.showwarning(
                "Stock insuficiente",
                f"No hay suficiente stock.\nStock disponible: {stock_actual}\nYa agregado en la venta: {cantidad_en_carrito}"
            )
            return

        subtotal = round(precio_final * cantidad_vendida, 2)

        carrito.append({
            "producto_id": producto_id,
            "nombre": nombre_producto,
            "cantidad": cantidad_vendida,
            "precio_normal": precio_normal,
            "descuento": descuento,
            "precio_final": precio_final,
            "subtotal": subtotal,
            "stock_actual": stock_actual
        })

        actualizar_tabla_carrito()

        combo_nombre.set("")
        combo_nombre["values"] = lista_productos
        entry_cantidad.delete(0, tk.END)
        label_info.config(text="")

    def quitar_seleccionado():
        seleccion = tabla_carrito.selection()

        if not seleccion:
            messagebox.showwarning("Aviso", "Selecciona un producto de la lista para quitar")
            return

        indice = int(seleccion[0])
        del carrito[indice]
        actualizar_tabla_carrito()

    def guardar_toda_la_venta():
        if not carrito:
            return 0, 0

        conexion, cursor = conectar()
        fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        total_general = 0
        stock_bajo_detectado = 0

        for item in carrito:
            producto_id = item["producto_id"]
            nombre = item["nombre"]
            cantidad = item["cantidad"]
            subtotal = item["subtotal"]

            cursor.execute("SELECT stock FROM productos WHERE id = ?", (producto_id,))
            fila_stock = cursor.fetchone()

            if fila_stock is None:
                continue

            stock_actual = fila_stock[0]

            if cantidad > stock_actual:
                conexion.close()
                messagebox.showerror(
                    "Error",
                    f"Stock insuficiente para {nombre}.\nStock actual: {stock_actual}"
                )
                return None, None

            nuevo_stock = stock_actual - cantidad

            cursor.execute(
                "UPDATE productos SET stock = ? WHERE id = ?",
                (nuevo_stock, producto_id)
            )

            cursor.execute(
                "INSERT INTO ventas (nombre_producto, cantidad, total, fecha) VALUES (?, ?, ?, ?)",
                (nombre, cantidad, subtotal, fecha)
            )

            total_general += subtotal

            if nuevo_stock <= 10:
                stock_bajo_detectado += 1

        conexion.commit()
        conexion.close()

        try:
            cantidad_ventas_hoy, total_vendido_hoy = obtener_ventas_hoy()
            label_ventas_hoy.config(text=f"Ventas realizadas: {cantidad_ventas_hoy}")
            label_total_hoy.config(text=f"Total vendido: {total_vendido_hoy} Bs")
        except:
            pass

        carrito.clear()
        actualizar_tabla_carrito()
        combo_nombre.set("")
        combo_nombre["values"] = lista_productos
        entry_cantidad.delete(0, tk.END)
        label_info.config(text="")

        return round(total_general, 2), stock_bajo_detectado

    def mostrar_qr():
        ventana_qr = tk.Toplevel(ventana)
        ventana_qr.title("Pago por QR")
        ventana_qr.geometry("380x500")
        ventana_qr.configure(bg="#f4f6f8")
        ventana_qr.resizable(False, False)

        total_general = sum(item["subtotal"] for item in carrito)

        tk.Label(
            ventana_qr,
            text="PAGO CON QR",
            font=("Arial", 16, "bold"),
            bg="#f4f6f8",
            fg="#1f3b5c"
        ).pack(pady=15)

        tk.Label(
            ventana_qr,
            text=f"Total a pagar: {round(total_general, 2)} Bs",
            font=("Arial", 14, "bold"),
            bg="#f4f6f8",
            fg="#27ae60"
        ).pack(pady=8)

        tk.Label(
            ventana_qr,
            text="Escanee el QR y realice el pago",
            font=("Arial", 10),
            bg="#f4f6f8",
            fg="#555"
        ).pack(pady=5)

        try:
            ruta_qr = os.path.join(os.path.dirname(__file__), "qr_pago.png")
            print("Buscando QR en:", ruta_qr)

            imagen = Image.open(ruta_qr)
            imagen = imagen.resize((220, 220))
            imagen_qr = ImageTk.PhotoImage(imagen)

            label_qr = tk.Label(ventana_qr, image=imagen_qr, bg="#f4f6f8")
            label_qr.image = imagen_qr
            label_qr.pack(pady=10)

        except Exception as e:
                print("ERROR QR:", e)
                tk.Label(
                    ventana_qr,
                    text=f"No se pudo cargar el QR\n{e}",
                    font=("Arial", 10, "italic"),
                    bg="#f4f6f8",
                    fg="red"
                ).pack(pady=40)
        except:
            tk.Label(
                ventana_qr,
                text="[No se encontró la imagen qr_pago.png]",
                font=("Arial", 10, "italic"),
                bg="#f4f6f8",
                fg="red"
            ).pack(pady=40)

        def confirmar_pago():
            total_guardado, stock_bajo_detectado = guardar_toda_la_venta()

            if total_guardado is None:
                return

            ventana_qr.destroy()

            mensaje = (
                f"Gracias por su compra\n\n"
                f"Total pagado: {total_guardado} Bs\n"
                f"Productos vendidos: {len(carrito_backup)}"
            )

            if stock_bajo_detectado > 0:
                mensaje += f"\n\n⚠ {stock_bajo_detectado} producto(s) quedaron con stock bajo"

            messagebox.showinfo("Pago recibido", mensaje)

        carrito_backup = carrito.copy()

        tk.Button(
            ventana_qr,
            text="Pago recibido",
            width=18,
            height=2,
            bg="#27ae60",
            fg="white",
            font=("Arial", 11, "bold"),
            command=confirmar_pago
        ).pack(pady=20)

    def cobrar():
        if not carrito:
            messagebox.showwarning("Aviso", "No hay productos agregados en la venta")
            return

        total_general = round(sum(item["subtotal"] for item in carrito), 2)

        ventana_confirmacion = tk.Toplevel(ventana)
        ventana_confirmacion.title("Confirmar compra")
        ventana_confirmacion.geometry("430x450")
        ventana_confirmacion.configure(bg="#f8f9fa")
        ventana_confirmacion.resizable(False, False)

        tk.Label(
            ventana_confirmacion,
            text="CONFIRMAR COMPRA",
            font=("Arial", 16, "bold"),
            bg="#f8f9fa",
            fg="#1f3b5c"
        ).pack(pady=15)

        frame_info = tk.Frame(ventana_confirmacion, bg="white", bd=1, relief="solid")
        frame_info.pack(padx=20, pady=10, fill="x")

        tk.Label(
            frame_info,
            text=f"Productos en la venta: {len(carrito)}",
            font=("Arial", 11, "bold"),
            bg="white",
            anchor="w"
        ).pack(fill="x", padx=15, pady=8)

        for item in carrito:
            tk.Label(
                frame_info,
                text=f'- {item["nombre"]} x{item["cantidad"]} = {item["subtotal"]} Bs',
                font=("Arial", 10),
                bg="white",
                anchor="w"
            ).pack(fill="x", padx=15, pady=2)

        tk.Label(
            frame_info,
            text=f"TOTAL: {total_general} Bs",
            font=("Arial", 13, "bold"),
            bg="white",
            fg="#27ae60",
            anchor="w"
        ).pack(fill="x", padx=15, pady=12)

        tk.Label(
            ventana_confirmacion,
            text="Método de pago",
            font=("Arial", 12, "bold"),
            bg="#f8f9fa",
            fg="#1f3b5c"
        ).pack(pady=(15, 8))

        metodo_pago = tk.StringVar(value="efectivo")

        frame_pago = tk.Frame(ventana_confirmacion, bg="#f8f9fa")
        frame_pago.pack()

        tk.Radiobutton(
            frame_pago,
            text="Efectivo",
            variable=metodo_pago,
            value="efectivo",
            font=("Arial", 11),
            bg="#f8f9fa"
        ).pack(side="left", padx=20)

        tk.Radiobutton(
            frame_pago,
            text="QR",
            variable=metodo_pago,
            value="qr",
            font=("Arial", 11),
            bg="#f8f9fa"
        ).pack(side="left", padx=20)

        def confirmar_cobro():
            metodo = metodo_pago.get()

            if metodo == "efectivo":
                total_guardado, stock_bajo_detectado = guardar_toda_la_venta()

                if total_guardado is None:
                    return

                ventana_confirmacion.destroy()

                mensaje = (
                    f"Venta confirmada\n\n"
                    f"Total cobrado: {total_guardado} Bs\n"
                    f"Productos vendidos: {len(carrito_backup)}"
                )

                if stock_bajo_detectado > 0:
                    mensaje += f"\n\n⚠ {stock_bajo_detectado} producto(s) quedaron con stock bajo"

                messagebox.showinfo("Venta", mensaje)

            else:
                ventana_confirmacion.destroy()
                mostrar_qr()

        carrito_backup = carrito.copy()

        frame_botones_confirm = tk.Frame(ventana_confirmacion, bg="#f8f9fa")
        frame_botones_confirm.pack(pady=25)

        tk.Button(
            frame_botones_confirm,
            text="Confirmar",
            width=14,
            height=2,
            bg="#27ae60",
            fg="white",
            font=("Arial", 10, "bold"),
            command=confirmar_cobro
        ).pack(side="left", padx=10)

        tk.Button(
            frame_botones_confirm,
            text="Cancelar",
            width=14,
            height=2,
            bg="#c0392b",
            fg="white",
            font=("Arial", 10, "bold"),
            command=ventana_confirmacion.destroy
        ).pack(side="left", padx=10)

    ventana = tk.Toplevel()
    ventana.title("Vender producto")
    ventana.geometry("900x620")
    ventana.configure(bg="#f4f6f8")
    ventana.resizable(False, False)

    titulo = tk.Label(
        ventana,
        text="VENTA DE PRODUCTOS",
        font=("Arial", 16, "bold"),
        bg="#f4f6f8",
        fg="#1f3b5c"
    )
    titulo.pack(pady=15)

    frame_form = tk.Frame(ventana, bg="#f4f6f8")
    frame_form.pack(pady=10)

    label_nombre = tk.Label(
        frame_form,
        text="Medicamento:",
        font=("Arial", 11),
        bg="#f4f6f8"
    )
    label_nombre.grid(row=0, column=0, padx=10, pady=10, sticky="e")

    combo_nombre = ttk.Combobox(
        frame_form,
        values=lista_productos,
        font=("Arial", 11),
        width=32
    )
    combo_nombre.grid(row=0, column=1, padx=10, pady=10)
    combo_nombre.bind("<KeyRelease>", filtrar_productos)

    label_cantidad = tk.Label(
        frame_form,
        text="Cantidad:",
        font=("Arial", 11),
        bg="#f4f6f8"
    )
    label_cantidad.grid(row=0, column=2, padx=10, pady=10, sticky="e")

    entry_cantidad = tk.Entry(frame_form, font=("Arial", 11), width=12)
    entry_cantidad.grid(row=0, column=3, padx=10, pady=10)

    frame_botones_superior = tk.Frame(ventana, bg="#f4f6f8")
    frame_botones_superior.pack(pady=10)

    btn_buscar = tk.Button(
        frame_botones_superior,
        text="Buscar",
        width=14,
        height=2,
        bg="#3498db",
        fg="white",
        font=("Arial", 10, "bold"),
        command=buscar_producto
    )
    btn_buscar.grid(row=0, column=0, padx=10)

    btn_agregar = tk.Button(
        frame_botones_superior,
        text="Agregar a la venta",
        width=18,
        height=2,
        bg="#27ae60",
        fg="white",
        font=("Arial", 10, "bold"),
        command=agregar_al_carrito
    )
    btn_agregar.grid(row=0, column=1, padx=10)

    btn_quitar = tk.Button(
        frame_botones_superior,
        text="Quitar seleccionado",
        width=18,
        height=2,
        bg="#c0392b",
        fg="white",
        font=("Arial", 10, "bold"),
        command=quitar_seleccionado
    )
    btn_quitar.grid(row=0, column=2, padx=10)

    label_info = tk.Label(
        ventana,
        text="",
        font=("Arial", 11, "bold"),
        bg="#f4f6f8",
        fg="#1f3b5c",
        wraplength=750,
        justify="center"
    )
    label_info.pack(pady=10)

    frame_tabla = tk.Frame(ventana, bg="#f4f6f8")
    frame_tabla.pack(fill="both", expand=True, padx=20, pady=10)

    columnas = ("Producto", "Cantidad", "Precio normal", "Descuento", "Precio final", "Subtotal")
    tabla_carrito = ttk.Treeview(frame_tabla, columns=columnas, show="headings", height=12)

    for col in columnas:
        tabla_carrito.heading(col, text=col)

    tabla_carrito.column("Producto", width=240, anchor="center")
    tabla_carrito.column("Cantidad", width=80, anchor="center")
    tabla_carrito.column("Precio normal", width=100, anchor="center")
    tabla_carrito.column("Descuento", width=90, anchor="center")
    tabla_carrito.column("Precio final", width=100, anchor="center")
    tabla_carrito.column("Subtotal", width=100, anchor="center")

    scrollbar = ttk.Scrollbar(frame_tabla, orient="vertical", command=tabla_carrito.yview)
    tabla_carrito.configure(yscrollcommand=scrollbar.set)

    tabla_carrito.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    label_total = tk.Label(
        ventana,
        text="TOTAL GENERAL: 0 Bs",
        font=("Arial", 15, "bold"),
        bg="#f4f6f8",
        fg="#27ae60"
    )
    label_total.pack(pady=10)

    btn_cobrar = tk.Button(
        ventana,
        text="Cobrar",
        width=20,
        height=2,
        bg="#8e44ad",
        fg="white",
        font=("Arial", 11, "bold"),
        command=cobrar
    )
    btn_cobrar.pack(pady=15)
def ver_historial():
    conexion, cursor = conectar()
    cursor.execute("SELECT * FROM ventas")
    ventas = cursor.fetchall()
    conexion.close()

    ventana = tk.Toplevel()
    ventana.title("Historial de ventas")
    ventana.geometry("800x400")
    ventana.configure(bg="#f4f6f8")

    titulo = tk.Label(
        ventana,
        text="HISTORIAL DE VENTAS",
        font=("Arial", 16, "bold"),
        bg="#f4f6f8",
        fg="#1f3b5c"
    )
    titulo.pack(pady=10)

    frame = tk.Frame(ventana, bg="#f4f6f8")
    frame.pack(fill="both", expand=True, padx=20, pady=10)

    columnas = ("ID", "Medicamento", "Cantidad", "Total", "Fecha")

    tabla = ttk.Treeview(frame, columns=columnas, show="headings", height=12)
    tabla.heading("ID", text="ID")
    tabla.heading("Medicamento", text="Medicamento")
    tabla.heading("Cantidad", text="Cantidad")
    tabla.heading("Total", text="Total")
    tabla.heading("Fecha", text="Fecha")

    tabla.column("ID", width=60, anchor="center")
    tabla.column("Medicamento", width=180, anchor="center")
    tabla.column("Cantidad", width=100, anchor="center")
    tabla.column("Total", width=100, anchor="center")
    tabla.column("Fecha", width=250, anchor="center")

    scrollbar = ttk.Scrollbar(frame, orient="vertical", command=tabla.yview)
    tabla.configure(yscrollcommand=scrollbar.set)

    tabla.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    for venta in ventas:
        tabla.insert("", "end", values=(venta[0], venta[1], venta[2], venta[3], venta[4]))
def obtener_ventas_hoy():
    conexion, cursor = conectar()

    hoy = datetime.now().strftime("%Y-%m-%d")

    cursor.execute(
        "SELECT COUNT(*), COALESCE(SUM(total), 0) FROM ventas WHERE fecha LIKE ?",
        (hoy + "%",)
    )

    resultado = cursor.fetchone()
    conexion.close()

    cantidad_ventas = resultado[0]
    total_vendido = resultado[1]

    return cantidad_ventas, total_vendido
def agregar_producto():

    ventana_agregar = tk.Toplevel()
    ventana_agregar.title("Agregar producto")
    ventana_agregar.geometry("350x320")

    tk.Label(ventana_agregar, text="Código:").pack(pady=5)
    entry_codigo = tk.Entry(ventana_agregar, width=30)
    entry_codigo.pack()

    tk.Label(ventana_agregar, text="Producto:").pack(pady=5)
    entry_nombre = tk.Entry(ventana_agregar, width=30)
    entry_nombre.pack()

    tk.Label(ventana_agregar, text="Marca:").pack(pady=5)
    entry_marca = tk.Entry(ventana_agregar, width=30)
    entry_marca.pack()

    tk.Label(ventana_agregar, text="Precio:").pack(pady=5)
    entry_precio = tk.Entry(ventana_agregar, width=30)
    entry_precio.pack()

    tk.Label(ventana_agregar, text="Stock:").pack(pady=5)
    entry_stock = tk.Entry(ventana_agregar, width=30)
    entry_stock.pack()

    def guardar():

        codigo = entry_codigo.get().strip()
        nombre = entry_nombre.get().strip()
        marca = entry_marca.get().strip()

        try:
            precio = float(entry_precio.get())
            stock = int(entry_stock.get())
        except:
            messagebox.showerror("Error", "Precio o stock inválido")
            return

        conexion, cursor = conectar()

        cursor.execute(
            "INSERT INTO productos (codigo,nombre,marca,precio,stock) VALUES (?,?,?,?,?)",
            (codigo,nombre,marca,precio,stock)
        )

        conexion.commit()
        conexion.close()

        messagebox.showinfo("Éxito", "Producto agregado")

        ventana_agregar.destroy()

    tk.Button(
        ventana_agregar,
        text="Guardar producto",
        bg="#27ae60",
        fg="white",
        command=guardar
    ).pack(pady=15)
import tkinter as tk
from tkinter import ttk
from datetime import datetime

def obtener_alertas_vencimiento():
    conexion, cursor = conectar()
    cursor.execute("""
        SELECT codigo, nombre, marca, stock, vencimiento
        FROM productos
        WHERE vencimiento IS NOT NULL AND vencimiento != ''
    """)
    productos = cursor.fetchall()
    conexion.close()

    hoy = datetime.now().date()
    vencidos = []
    por_vencer = []

    for producto in productos:
        codigo, nombre, marca, stock, vencimiento = producto

        try:
            fecha_venc = datetime.strptime(vencimiento, "%Y-%m-%d").date()
            dias = (fecha_venc - hoy).days

            if dias < 0:
                vencidos.append((codigo, nombre, marca, stock, vencimiento, abs(dias)))
            elif dias <= 30:
                por_vencer.append((codigo, nombre, marca, stock, vencimiento, dias))
        except:
            pass

    return vencidos, por_vencer


def ver_alertas_vencimiento():
    vencidos, por_vencer = obtener_alertas_vencimiento()

    ventana = tk.Toplevel()
    ventana.title("Alertas de vencimiento")
    ventana.geometry("950x550")
    ventana.config(bg="#f2f2f2")

    # Título principal
    titulo = tk.Label(
        ventana,
        text="ALERTAS DE VENCIMIENTO",
        font=("Arial", 18, "bold"),
        bg="#f2f2f2",
        fg="#1f3b5c"
    )
    titulo.pack(pady=10)

    # =========================
    # SECCIÓN VENCIDOS
    # =========================
    frame_vencidos = tk.LabelFrame(
        ventana,
        text=f" Productos vencidos ({len(vencidos)}) ",
        font=("Arial", 12, "bold"),
        fg="red",
        bg="#f2f2f2",
        padx=10,
        pady=10
    )
    frame_vencidos.pack(fill="both", expand=True, padx=15, pady=10)

    columnas_vencidos = ("codigo", "nombre", "marca", "stock", "vencimiento", "dias")
    tabla_vencidos = ttk.Treeview(frame_vencidos, columns=columnas_vencidos, show="headings", height=8)

    tabla_vencidos.heading("codigo", text="Código")
    tabla_vencidos.heading("nombre", text="Producto")
    tabla_vencidos.heading("marca", text="Marca")
    tabla_vencidos.heading("stock", text="Stock")
    tabla_vencidos.heading("vencimiento", text="Fecha venc.")
    tabla_vencidos.heading("dias", text="Días vencido")

    tabla_vencidos.column("codigo", width=100, anchor="center")
    tabla_vencidos.column("nombre", width=280)
    tabla_vencidos.column("marca", width=180)
    tabla_vencidos.column("stock", width=80, anchor="center")
    tabla_vencidos.column("vencimiento", width=120, anchor="center")
    tabla_vencidos.column("dias", width=100, anchor="center")

    scroll1 = ttk.Scrollbar(frame_vencidos, orient="vertical", command=tabla_vencidos.yview)
    tabla_vencidos.configure(yscrollcommand=scroll1.set)

    tabla_vencidos.pack(side="left", fill="both", expand=True)
    scroll1.pack(side="right", fill="y")

    if vencidos:
        for p in vencidos:
            tabla_vencidos.insert("", tk.END, values=p)
    else:
        tabla_vencidos.insert("", tk.END, values=("", "No hay productos vencidos", "", "", "", ""))

    # =========================
    # SECCIÓN POR VENCER
    # =========================
    frame_por_vencer = tk.LabelFrame(
        ventana,
        text=f" Productos por vencer en 30 días ({len(por_vencer)}) ",
        font=("Arial", 12, "bold"),
        fg="#c98900",
        bg="#f2f2f2",
        padx=10,
        pady=10
    )
    frame_por_vencer.pack(fill="both", expand=True, padx=15, pady=10)

    columnas_por_vencer = ("codigo", "nombre", "marca", "stock", "vencimiento", "dias")
    tabla_por_vencer = ttk.Treeview(frame_por_vencer, columns=columnas_por_vencer, show="headings", height=8)

    tabla_por_vencer.heading("codigo", text="Código")
    tabla_por_vencer.heading("nombre", text="Producto")
    tabla_por_vencer.heading("marca", text="Marca")
    tabla_por_vencer.heading("stock", text="Stock")
    tabla_por_vencer.heading("vencimiento", text="Fecha venc.")
    tabla_por_vencer.heading("dias", text="Días restantes")

    tabla_por_vencer.column("codigo", width=100, anchor="center")
    tabla_por_vencer.column("nombre", width=280)
    tabla_por_vencer.column("marca", width=180)
    tabla_por_vencer.column("stock", width=80, anchor="center")
    tabla_por_vencer.column("vencimiento", width=120, anchor="center")
    tabla_por_vencer.column("dias", width=110, anchor="center")

    scroll2 = ttk.Scrollbar(frame_por_vencer, orient="vertical", command=tabla_por_vencer.yview)
    tabla_por_vencer.configure(yscrollcommand=scroll2.set)

    tabla_por_vencer.pack(side="left", fill="both", expand=True)
    scroll2.pack(side="right", fill="y")

    if por_vencer:
        for p in por_vencer:
            tabla_por_vencer.insert("", tk.END, values=p)
    else:
        tabla_por_vencer.insert("", tk.END, values=("", "No hay productos por vencer", "", "", "", ""))
def ver_promociones():
    ventana = tk.Toplevel()
    ventana.title("Promociones")
    ventana.geometry("900x520")
    ventana.configure(bg="#f4f6f8")

    tk.Label(
        ventana,
        text="GESTIÓN DE PROMOCIONES",
        font=("Arial", 16, "bold"),
        bg="#f4f6f8",
        fg="#1f3b5c"
    ).pack(pady=15)

    frame_form = tk.Frame(ventana, bg="#f4f6f8")
    frame_form.pack(pady=10)

    tk.Label(frame_form, text="Producto:", font=("Arial", 11), bg="#f4f6f8").grid(row=0, column=0, padx=8, pady=8, sticky="e")
    tk.Label(frame_form, text="Descuento (%):", font=("Arial", 11), bg="#f4f6f8").grid(row=1, column=0, padx=8, pady=8, sticky="e")
    tk.Label(frame_form, text="Fecha inicio:", font=("Arial", 11), bg="#f4f6f8").grid(row=2, column=0, padx=8, pady=8, sticky="e")
    tk.Label(frame_form, text="Fecha fin:", font=("Arial", 11), bg="#f4f6f8").grid(row=3, column=0, padx=8, pady=8, sticky="e")

    combo_producto = ttk.Combobox(frame_form, width=45)
    combo_producto.grid(row=0, column=1, padx=8, pady=8)
    lista_productos = []
    def filtrar_productos(event=None):
        texto = combo_producto.get().strip().lower()

        if texto == "":
            combo_producto["values"] = lista_productos
        else:
            filtrados = [p for p in lista_productos if texto in p.lower()]
            combo_producto["values"] = filtrados
    combo_producto.bind("<KeyRelease>", filtrar_productos)
    entry_descuento = tk.Entry(frame_form, width=20)
    entry_descuento.grid(row=1, column=1, padx=8, pady=8, sticky="w")

    entry_inicio = tk.Entry(frame_form, width=20)
    entry_inicio.grid(row=2, column=1, padx=8, pady=8, sticky="w")

    entry_fin = tk.Entry(frame_form, width=20)
    entry_fin.grid(row=3, column=1, padx=8, pady=8, sticky="w")

    tk.Label(frame_form, text="Formato: YYYY-MM-DD", font=("Arial", 9), bg="#f4f6f8", fg="gray").grid(row=2, column=2, padx=5)
    tk.Label(frame_form, text="Formato: YYYY-MM-DD", font=("Arial", 9), bg="#f4f6f8", fg="gray").grid(row=3, column=2, padx=5)

    productos_dict = {}

    def cargar_productos():
        conexion, cursor = conectar()
        cursor.execute("SELECT id, nombre, marca, precio, stock FROM productos ORDER BY nombre")
        productos = cursor.fetchall()
        conexion.close()

        lista = []
        productos_dict.clear()

        for p in productos:
            producto_id, nombre, marca, precio, stock = p
            texto = f"{nombre} | Marca: {marca} | Precio: {precio} | Stock: {stock}"
            lista.append(texto)
            productos_dict[texto] = producto_id
        lista_productos.clear()
        lista_productos.extend(lista)
        combo_producto["values"] = lista

    def cargar_promociones():
        for item in tabla.get_children():
            tabla.delete(item)

        conexion, cursor = conectar()
        cursor.execute("""
            SELECT promociones.id, productos.nombre, productos.marca, promociones.descuento,
                   promociones.fecha_inicio, promociones.fecha_fin, promociones.activa
            FROM promociones
            INNER JOIN productos ON promociones.producto_id = productos.id
            ORDER BY promociones.id DESC
        """)
        promociones = cursor.fetchall()
        conexion.close()

        hoy = datetime.now().date()

        for promo in promociones:
            promo_id, nombre, marca, descuento, fecha_inicio, fecha_fin, activa = promo

            estado = "Inactiva"
            try:
                inicio = datetime.strptime(fecha_inicio, "%Y-%m-%d").date()
                fin = datetime.strptime(fecha_fin, "%Y-%m-%d").date()

                if activa == 1:
                    if inicio <= hoy <= fin:
                        estado = "Activa"
                    elif hoy < inicio:
                        estado = "Pendiente"
                    else:
                        estado = "Finalizada"
            except:
                estado = "Fecha inválida"

            tabla.insert("", "end", values=(
                promo_id, nombre, marca, f"{descuento}%", fecha_inicio, fecha_fin, estado
            ))

    def guardar_promocion():
        producto_texto = combo_producto.get().strip()
        descuento = entry_descuento.get().strip()
        fecha_inicio = entry_inicio.get().strip()
        fecha_fin = entry_fin.get().strip()

        if not producto_texto or not descuento or not fecha_inicio or not fecha_fin:
            messagebox.showwarning("Aviso", "Completa todos los campos")
            return

        try:
            descuento = float(descuento)
            if descuento <= 0 or descuento >= 100:
                messagebox.showwarning("Aviso", "El descuento debe ser mayor a 0 y menor a 100")
                return
        except:
            messagebox.showerror("Error", "El descuento debe ser un número")
            return

        try:
            inicio = datetime.strptime(fecha_inicio, "%Y-%m-%d")
            fin = datetime.strptime(fecha_fin, "%Y-%m-%d")
            if fin < inicio:
                messagebox.showwarning("Aviso", "La fecha fin no puede ser menor a la fecha inicio")
                return
        except:
            messagebox.showerror("Error", "Las fechas deben tener formato YYYY-MM-DD")
            return

        producto_id = productos_dict.get(producto_texto)
        if not producto_id:
            messagebox.showerror("Error", "Selecciona un producto válido")
            return

        conexion, cursor = conectar()
        cursor.execute("""
            INSERT INTO promociones (producto_id, descuento, fecha_inicio, fecha_fin, activa)
            VALUES (?, ?, ?, ?, 1)
        """, (producto_id, descuento, fecha_inicio, fecha_fin))
        conexion.commit()
        conexion.close()

        messagebox.showinfo("Éxito", "Promoción guardada correctamente")

        combo_producto.set("")
        entry_descuento.delete(0, tk.END)
        entry_inicio.delete(0, tk.END)
        entry_fin.delete(0, tk.END)

        cargar_promociones()

    def eliminar_promocion():
        seleccion = tabla.selection()
        if not seleccion:
            messagebox.showwarning("Aviso", "Selecciona una promoción")
            return

        datos = tabla.item(seleccion[0], "values")
        promo_id = datos[0]

        confirmar = messagebox.askyesno("Confirmar", "¿Eliminar esta promoción?")
        if not confirmar:
            return

        conexion, cursor = conectar()
        cursor.execute("DELETE FROM promociones WHERE id = ?", (promo_id,))
        conexion.commit()
        conexion.close()

        cargar_promociones()
        messagebox.showinfo("Éxito", "Promoción eliminada")

    def activar_promocion():
        seleccion = tabla.selection()
        if not seleccion:
            messagebox.showwarning("Aviso", "Selecciona una promoción")
            return

        datos = tabla.item(seleccion[0], "values")
        promo_id = datos[0]

        conexion, cursor = conectar()
        cursor.execute("UPDATE promociones SET activa = 1 WHERE id = ?", (promo_id,))
        conexion.commit()
        conexion.close()

        cargar_promociones()
        messagebox.showinfo("Éxito", "Promoción activada")

    def desactivar_promocion():
        seleccion = tabla.selection()
        if not seleccion:
            messagebox.showwarning("Aviso", "Selecciona una promoción")
            return

        datos = tabla.item(seleccion[0], "values")
        promo_id = datos[0]

        conexion, cursor = conectar()
        cursor.execute("UPDATE promociones SET activa = 0 WHERE id = ?", (promo_id,))
        conexion.commit()
        conexion.close()

        cargar_promociones()
        messagebox.showinfo("Éxito", "Promoción desactivada")

    tk.Button(
        frame_form,
        text="Guardar promoción",
        bg="#27ae60",
        fg="white",
        font=("Arial", 10, "bold"),
        width=18,
        command=guardar_promocion
    ).grid(row=4, column=0, columnspan=2, pady=12)

    frame_tabla = tk.Frame(ventana, bg="#f4f6f8")
    frame_tabla.pack(fill="both", expand=True, padx=20, pady=10)

    columnas = ("ID", "Producto", "Marca", "Descuento", "Inicio", "Fin", "Estado")
    tabla = ttk.Treeview(frame_tabla, columns=columnas, show="headings", height=12)

    for col in columnas:
        tabla.heading(col, text=col)

    tabla.column("ID", width=50, anchor="center")
    tabla.column("Producto", width=220, anchor="center")
    tabla.column("Marca", width=150, anchor="center")
    tabla.column("Descuento", width=90, anchor="center")
    tabla.column("Inicio", width=100, anchor="center")
    tabla.column("Fin", width=100, anchor="center")
    tabla.column("Estado", width=100, anchor="center")

    scrollbar = ttk.Scrollbar(frame_tabla, orient="vertical", command=tabla.yview)
    tabla.configure(yscrollcommand=scrollbar.set)

    tabla.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    frame_botones = tk.Frame(ventana, bg="#f4f6f8")
    frame_botones.pack(pady=10)

    tk.Button(frame_botones, text="Activar", width=12, bg="#3498db", fg="white", command=activar_promocion).pack(side="left", padx=5)
    tk.Button(frame_botones, text="Desactivar", width=12, bg="#f39c12", fg="white", command=desactivar_promocion).pack(side="left", padx=5)
    tk.Button(frame_botones, text="Eliminar", width=12, bg="#c0392b", fg="white", command=eliminar_promocion).pack(side="left", padx=5)

    cargar_productos()
    cargar_promociones()
ventana_principal = tk.Tk()
ventana_principal.title("Sistema de Farmacia")
ventana_principal.geometry("850x760")
ventana_principal.configure(bg="#eaf2f8")
ventana_principal.resizable(False, False)
cantidad_ventas_hoy, total_vendido_hoy = obtener_ventas_hoy()

style = ttk.Style()
style.theme_use("clam")

titulo = tk.Label(
    ventana_principal,
    text="SISTEMA DE FARMACIA",
    font=("Arial", 20, "bold"),
    bg="#eaf2f8",
    fg="#1f3b5c"
)
titulo.pack(pady=25)

subtitulo = tk.Label(
    ventana_principal,
    text="Control de inventario y ventas",
    font=("Arial", 11),
    bg="#eaf2f8",
    fg="#4a6572"
)
subtitulo.pack(pady=5)
frame_resumen = tk.Frame(ventana_principal, bg="#d6eaf8", bd=1, relief="solid")
frame_resumen.pack(pady=10, padx=40, fill="x")

label_resumen = tk.Label(
    frame_resumen,
    text="RESUMEN DE HOY",
    font=("Arial", 11, "bold"),
    bg="#d6eaf8",
    fg="#1f3b5c"
)
label_resumen.pack(pady=(6,2))

label_ventas_hoy = tk.Label(
    frame_resumen,
    text=f"Ventas realizadas: {cantidad_ventas_hoy}",
    font=("Arial", 10, "bold"),
    bg="#d6eaf8",
    fg="#1f3b5c"
)
label_ventas_hoy.pack()

label_total_hoy = tk.Label(
    frame_resumen,
    text=f"Total vendido: {total_vendido_hoy} Bs",
    font=("Arial", 10, "bold"),
    bg="#d6eaf8",
    fg="#1f3b5c"
)
label_total_hoy.pack(pady=(0,6))

frame_botones = tk.Frame(ventana_principal, bg="#eaf2f8")
frame_botones.pack(pady=20)

btn_inventario = tk.Button(
    frame_botones, text="Ver inventario", width=25, height=2,
    bg="#28b463", fg="white", font=("Arial", 11, "bold"),
    command=ver_inventario
)
btn_inventario.pack(pady=12)

btn_vender = tk.Button(
    frame_botones, text="Vender producto", width=25, height=2,
    bg="#f39c12", fg="white", font=("Arial", 11, "bold"),
    command=vender_producto
)
btn_vender.pack(pady=12)

btn_historial = tk.Button(
    frame_botones, text="Ver historial de ventas", width=25, height=2,
    bg="#8e44ad", fg="white", font=("Arial", 11, "bold"),
    command=ver_historial
)
btn_historial.pack(pady=12)
btn_alertas = tk.Button(
    frame_botones,
    text="⚠ Alertas de vencimiento",
    width=25,
    height=2,
    bg="#c0392b",
    fg="white",
    font=("Arial", 11, "bold"),
    command=ver_alertas_vencimiento
)
btn_alertas.pack(pady=12)
btn_promociones = tk.Button(
    frame_botones,
    text="Promociones",
    width=25,
    height=2,
    bg="#2980b9",
    fg="white",
    font=("Arial", 11, "bold"),
    command=ver_promociones
)
btn_promociones.pack(pady=12)

btn_salir = tk.Button(
    frame_botones, text="Salir", width=25, height=2,
    bg="#c0392b", fg="white", font=("Arial", 11, "bold"),
    command=ventana_principal.quit
)
btn_salir.pack(pady=12)

ventana_principal.mainloop()